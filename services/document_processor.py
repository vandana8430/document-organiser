"""Document validation, extraction, and embedding helpers."""
import os, re, zlib, subprocess, io, zipfile, tempfile, shutil, time
from collections import Counter
import httpx
import numpy as np
# ---------- CLIP: image understanding (auto-tags + visual search) ----------
_c = None
def clip():
    global _c
    if _c is None:
        try:
            from sentence_transformers import SentenceTransformer
            _c = SentenceTransformer("clip-ViT-B-32")
        except Exception:
            _c = False
    return _c

LABELS = ["photo of a person", "group photo", "certificate", "screenshot", "diagram or chart", "landscape", "food", "animal",
          "building", "vehicle", "id card", "laptop or computer", "presentation slide", "handwritten note", "document page"]

def image_info(path):
    c = clip()
    if not c: return None, []
    from PIL import Image
    e = c.encode(Image.open(path), normalize_embeddings=True)
    s = c.encode(LABELS, normalize_embeddings=True) @ e
    return e.astype(np.float32), [LABELS[i] for i in np.argsort(-s)[:2] if s[i] > 0.2]

# ---------- Embeddings (Ollama embeddinggemma, optional MiniLM, lexical fallback) ----------
_m = None
_active_embedding_backend = None
_ollama_embedding_retry_after = 0.0


def _embedding_settings():
    model = os.getenv("OLLAMA_EMBEDDING_MODEL", "embeddinggemma").strip() or "embeddinggemma"
    base_url = os.getenv("OLLAMA_BASE_URL", "http://localhost:11434").strip().rstrip("/")
    url = os.getenv("OLLAMA_EMBEDDING_URL", "").strip() or f"{base_url}/api/embed"
    return model, url


def _embed_with_ollama(texts, model, url):
    vectors = []
    timeout = float(os.getenv("OLLAMA_EMBED_TIMEOUT_SECONDS", "90"))
    timeout = min(300.0, max(5.0, timeout))
    request_timeout = httpx.Timeout(timeout, connect=min(4.0, timeout))
    for start in range(0, len(texts), 32):
        batch = texts[start:start + 32]
        response = httpx.post(url, json={"model": model, "input": batch}, timeout=request_timeout)
        response.raise_for_status()
        payload = response.json()
        batch_vectors = payload.get("embeddings") if isinstance(payload, dict) else None
        if batch_vectors is None and isinstance(payload, dict) and len(batch) == 1:
            one = payload.get("embedding")
            batch_vectors = [one] if isinstance(one, list) else None
        if not isinstance(batch_vectors, list) or len(batch_vectors) != len(batch):
            raise ValueError("Ollama returned an invalid embedding response")
        vectors.extend(batch_vectors)
    result = np.asarray(vectors, dtype=np.float32)
    if result.ndim != 2 or result.shape[0] != len(texts) or result.shape[1] == 0 or not np.isfinite(result).all():
        raise ValueError("Ollama returned invalid embedding vectors")
    return result / (np.linalg.norm(result, axis=1, keepdims=True) + 1e-9)


def embed(texts):
    global _m, _active_embedding_backend, _ollama_embedding_retry_after
    texts = [str(text or "") for text in texts]
    if not texts:
        return np.zeros((0, 384), np.float32)
    ollama_enabled = os.getenv("SMARTFILES_OLLAMA_EMBEDDINGS", "1").strip().lower() not in {"0", "false", "no"}
    if ollama_enabled and time.monotonic() >= _ollama_embedding_retry_after:
        try:
            model, url = _embedding_settings()
            result = _embed_with_ollama(texts, model, url)
            _active_embedding_backend = "ollama:" + model
            _ollama_embedding_retry_after = 0.0
            return result
        except Exception:
            # Keep uploads and local search usable when Ollama is temporarily
            # unavailable; the bounded retry avoids one failed request per chunk.
            _ollama_embedding_retry_after = time.monotonic() + 30.0
    if _m is None:
        try:
            from sentence_transformers import SentenceTransformer
            _m = SentenceTransformer("all-MiniLM-L6-v2")
        except Exception:
            _m = False
    if _m:
        _active_embedding_backend = "minilm"
        return np.asarray(_m.encode(texts, normalize_embeddings=True), dtype=np.float32)
    out = np.zeros((len(texts), 384), np.float32)
    for i, t in enumerate(texts):
        for w in re.findall(r"\w+", t.lower()):
            out[i, zlib.crc32(w.encode()) % 384] += 1
    _active_embedding_backend = "hashed-v1"
    return out / (np.linalg.norm(out, axis=1, keepdims=True) + 1e-9)

blob = lambda v: v.astype(np.float32).tobytes()
unblob = lambda b: np.frombuffer(b, np.float32)

def embedding_model_state():
    """Return whether the active vectors use a semantic model, if initialized."""
    if _active_embedding_backend is not None:
        return _active_embedding_backend != "hashed-v1"
    return None if _m is None else _m is not False

def embedding_backend():
    """Return a durable model label so incompatible vectors are never compared."""
    if _active_embedding_backend is not None:
        return _active_embedding_backend
    return "minilm" if _m is not None and _m is not False else "hashed-v1"

# ---------- Content extraction (docs, OCR, video transcripts) ----------
KINDS = {"document": {".pdf", ".docx", ".txt", ".md", ".csv", ".xlsx", ".pptx"},
         "image": {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"},
         "video": {".mp4", ".mkv", ".mov", ".avi", ".webm"}}
MAX_EXTRACTED_CHARS = int(os.getenv("SMARTFILES_MAX_EXTRACTED_CHARS", "25000000"))

def _configure_tesseract(pytesseract):
    """Find the common Windows installer location even before PATH is refreshed."""
    configured = os.getenv("TESSERACT_CMD", "").strip()
    if configured and os.path.isfile(configured):
        pytesseract.pytesseract.tesseract_cmd = configured
        return
    current = pytesseract.pytesseract.tesseract_cmd
    if shutil.which(current):
        return
    for candidate in (
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ):
        if os.path.isfile(candidate):
            pytesseract.pytesseract.tesseract_cmd = candidate
            return

def _validate_ooxml(data, required_parts, description):
    """Validate an Office ZIP container without extracting archive paths."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 10000:
                raise ValueError(f"The {description} contains too many archive entries")
            expanded_size = 0
            for member in members:
                if member.flag_bits & 0x1:
                    raise ValueError(f"Encrypted {description} files are not supported")
                expanded_size += member.file_size
                if member.file_size > 50 * 1024 * 1024:
                    raise ValueError(f"The {description} contains an oversized component")
                if member.file_size / max(member.compress_size, 1) > 200:
                    raise ValueError(f"The {description} has a suspicious compression ratio")
                if expanded_size > 120 * 1024 * 1024:
                    raise ValueError(f"The {description} expands beyond the safe processing limit")
            names = set(archive.namelist())
            if not set(required_parts).issubset(names):
                raise ValueError(f"The file is not a valid {description}")
            # Read every member to detect corrupt or truncated ZIP entries.
            if archive.testzip() is not None:
                raise ValueError(f"The {description} archive is damaged")
    except (zipfile.BadZipFile, OSError, EOFError, NotImplementedError, RuntimeError):
        raise ValueError(f"The file is not a valid {description}")

def kind_of(ext):
    return next((k for k, s in KINDS.items() if ext in s), None)

def validate_upload_content(data, ext):
    """Reject common extension/content mismatches before storing an upload."""
    signatures = {
        ".png": lambda b: b.startswith(b"\x89PNG\r\n\x1a\n"),
        ".jpg": lambda b: b.startswith(b"\xff\xd8\xff"),
        ".jpeg": lambda b: b.startswith(b"\xff\xd8\xff"),
        ".webp": lambda b: len(b) >= 12 and b[:4] == b"RIFF" and b[8:12] == b"WEBP",
        ".bmp": lambda b: b.startswith(b"BM"),
        ".gif": lambda b: b.startswith((b"GIF87a", b"GIF89a")),
    }
    if ext == ".pdf" and b"%PDF-" not in data[:1024]:
        raise ValueError("The file does not contain a valid PDF header")
    if ext == ".docx":
        _validate_ooxml(data, ("[Content_Types].xml", "word/document.xml"), "DOCX document")
    if ext == ".xlsx":
        _validate_ooxml(data, ("[Content_Types].xml", "xl/workbook.xml"), "Excel workbook")
    if ext == ".pptx":
        _validate_ooxml(data, ("[Content_Types].xml", "ppt/presentation.xml"), "PowerPoint presentation")
    if ext in signatures and not signatures[ext](data):
        raise ValueError("The file contents do not match the selected image type")
    if ext in signatures:
        try:
            from PIL import Image
            with Image.open(io.BytesIO(data)) as image:
                if image.width * image.height > 50_000_000:
                    raise ValueError("The image is too large to process safely (maximum 50 megapixels)")
                image.verify()
        except ValueError:
            raise
        except Exception:
            raise ValueError("The image file is damaged or unreadable")
    if ext in {".md", ".csv"} and b"\x00" in data:
        raise ValueError("The file appears to be binary, not a text document")

def extract(path, kind, ext):
    """Returns (text, segments). segments = [(start_seconds, text)] for video."""
    if ext == ".pdf":
        from pypdf import PdfReader
        parts, size = [], 0
        for page_number, page in enumerate(PdfReader(path).pages):
            part = page.extract_text() or ""
            if not part.strip():
                # Scanned PDFs have no text layer. OCR only blank pages, keeping
                # ordinary PDFs fast and preserving page context in the result.
                import fitz
                import pytesseract
                from PIL import Image
                _configure_tesseract(pytesseract)
                with fitz.open(path) as pdf:
                    raster = pdf[page_number].get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    part = pytesseract.image_to_string(Image.open(io.BytesIO(raster.tobytes("png")))) or ""
            if part.strip():
                parts.append(f"Page {page_number + 1}\n{part}")
            size += len(part)
            if size >= MAX_EXTRACTED_CHARS: break
        return "\n".join(parts), None
    if ext == ".docx":
        from docx import Document
        document = Document(path)
        content, size = [], 0

        def table_lines(table):
            nonlocal size
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    values = [paragraph.text.strip() for paragraph in cell.paragraphs if paragraph.text.strip()]
                    for nested_table in cell.tables:
                        values.extend(table_lines(nested_table))
                    cells.append(" / ".join(values))
                line = " | ".join(cells)
                if line.strip():
                    content.append(line)
                    size += len(line)
                if size >= MAX_EXTRACTED_CHARS:
                    return

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                style_name = getattr(getattr(paragraph, "style", None), "name", "") or ""
                if style_name.lower().startswith("heading"):
                    content.append("Section: " + paragraph.text.strip())
                else:
                    content.append(paragraph.text)
                size += len(paragraph.text)
                if size >= MAX_EXTRACTED_CHARS: break
        if size < MAX_EXTRACTED_CHARS:
            for section in document.sections:
                for story in (section.header, section.footer):
                    for paragraph in story.paragraphs:
                        if paragraph.text.strip():
                            content.append(paragraph.text)
                            size += len(paragraph.text)
                            if size >= MAX_EXTRACTED_CHARS: break
                    if size >= MAX_EXTRACTED_CHARS: break
                    for table in story.tables:
                        table_lines(table)
                        if size >= MAX_EXTRACTED_CHARS: break
                    if size >= MAX_EXTRACTED_CHARS: break
                if size >= MAX_EXTRACTED_CHARS: break
        if size < MAX_EXTRACTED_CHARS:
            for table in document.tables:
                table_lines(table)
                if size >= MAX_EXTRACTED_CHARS: break
        return "\n".join(content), None
    if ext == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        lines, size, cells = [], 0, 0
        try:
            for sheet in workbook.worksheets[:100]:
                lines.append(f"Sheet: {sheet.title}")
                rows = sheet.iter_rows(max_row=min(sheet.max_row or 1, 10000),
                                       max_col=min(sheet.max_column or 1, 100), values_only=True)
                for values in rows:
                    cells += len(values)
                    rendered = [str(value) for value in values if value is not None]
                    if rendered:
                        line = " | ".join(rendered)
                        lines.append(line)
                        size += len(line)
                    if cells >= 50000 or size >= MAX_EXTRACTED_CHARS: break
                if cells >= 50000 or size >= MAX_EXTRACTED_CHARS: break
        finally:
            workbook.close()
        return "\n".join(lines), None
    if ext == ".pptx":
        from pptx import Presentation
        presentation = Presentation(path)
        lines, size = [], 0
        for number, slide in enumerate(presentation.slides):
            if number >= 500: break
            lines.append(f"Slide {number + 1}")
            for shape_number, shape in enumerate(slide.shapes):
                if shape_number >= 10000: break
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    lines.append(shape.text.strip())
                    size += len(shape.text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        line = " | ".join(cell.text for cell in row.cells)
                        lines.append(line)
                        size += len(line)
                if size >= MAX_EXTRACTED_CHARS: break
            if size >= MAX_EXTRACTED_CHARS: break
        return "\n".join(lines), None
    if kind == "image":  # OCR
        import pytesseract
        from PIL import Image
        _configure_tesseract(pytesseract)
        return pytesseract.image_to_string(Image.open(path)), None
    if kind == "video":  # FFmpeg audio extraction + speech-to-text
        with tempfile.NamedTemporaryFile(suffix=".wav", delete=False) as temporary_audio:
            wav = temporary_audio.name
        try:
            subprocess.run(["ffmpeg", "-y", "-i", path, "-vn", "-ac", "1", "-ar", "16000", wav],
                           capture_output=True, check=True)
            from faster_whisper import WhisperModel
            segs = [(s.start, s.text.strip()) for s in WhisperModel("base", compute_type="int8").transcribe(wav)[0]]
            return " ".join(t for _, t in segs), segs
        finally:
            try: os.remove(wav)
            except OSError: pass
    with open(path, "rb") as source:
        data = source.read(MAX_EXTRACTED_CHARS * 4 + 8)
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")[:MAX_EXTRACTED_CHARS], None
    return data.decode("utf-8-sig", errors="replace")[:MAX_EXTRACTED_CHARS], None

def make_chunks(text, segs):
    if segs:  # group ~5 transcript segments, keep start time for video moment retrieval
        return [(segs[i][0], " ".join(t for _, t in segs[i:i + 5])) for i in range(0, len(segs), 5)]
    text = text or ""
    marker = re.compile(r"(?im)^\s*(Page\s+\d+|Slide\s+\d+|Section:\s*[^\r\n]+|#{1,6}\s+[^\r\n]+)\s*$")
    markers = list(marker.finditer(text))
    if not markers:
        return [(i, text[i:i + 700]) for i in range(0, len(text), 600) if text[i:i + 700].strip()]

    chunks = []
    def add_segment(reference, body, offset):
        body = body.strip()
        if not body:
            return
        prefix = (reference + "\n") if reference else ""
        width = max(1, 700 - len(prefix))
        stride = max(1, min(600, width - 100))
        for position in range(0, len(body), stride):
            passage = prefix + body[position:position + width]
            if passage.strip():
                chunks.append((offset + position, passage))

    add_segment("", text[:markers[0].start()], 0)
    for index, match in enumerate(markers):
        end = markers[index + 1].start() if index + 1 < len(markers) else len(text)
        raw_reference = match.group(1).strip()
        reference = "Section: " + raw_reference.lstrip("# ").strip() if raw_reference.startswith("#") else raw_reference
        add_segment(reference, text[match.end():end], match.start())
    return chunks

