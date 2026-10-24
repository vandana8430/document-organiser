"""Integration tests against an isolated temporary database and upload folder."""
import os
import tempfile
import csv
import hashlib
import sqlite3
import pytest
from pathlib import Path

_test_root = tempfile.TemporaryDirectory(prefix="smartfiles-tests-")
os.environ["SMARTFILES_DB_PATH"] = str(Path(_test_root.name) / "test.db")
os.environ["SMARTFILES_UPLOAD_DIR"] = str(Path(_test_root.name) / "uploads")

from fastapi.testclient import TestClient
from app import app
from services import document_processor

# Keep integration tests local and deterministic; no model downloads or paid API calls.
document_processor._m = False
document_processor._c = False

client = TestClient(app)


def test_rule_categories_and_filename_fallback_tags():
    from services.categorizer import CATEGORIES, classify, keywords

    cases = [
        ("Asha_Mehta_Resume.pdf", "Professional experience and technical skills", "document", "Career & Resumes"),
        ("database_assignment.docx", "Semester coursework and university marks", "document", "Education & Assignments"),
        ("course_certificate.pdf", "Awarded to Mira Rao", "document", "Certificates"),
        ("project_report.pdf", "Project report with methodology and implementation details", "document", "Project Reports"),
        ("technical_notes.txt", "Python API database configuration notes", "document", "Technical Notes"),
        ("expense_data.csv", "Expense budget payment totals", "document", "Financial Documents"),
        ("launch.pptx", "", "document", "Presentations"),
        ("diagram.png", "", "image", "Images & Scanned Documents"),
        ("misc.pdf", "", "document", "Other Documents"),
    ]
    for name, text, kind, expected in cases:
        assert classify(text, kind, name) == expected
        assert expected in CATEGORIES
    assert "Python" in keywords("Python API database configuration", filename="technical_notes.txt")
    assert keywords("", filename="field_notes.txt")


def test_extraction_module_error_keeps_uploaded_file(monkeypatch):
    from services import document_ingestion
    from services import rag_assistant

    assert client.post("/api/register", json={"email": "extract-warning@example.test", "password": "extract-test-123"}).status_code == 200
    pdf = Path(__file__).resolve().parents[1] / "sample_documents" / "pdf" / "02_sample_project_report.pdf"

    def missing_pypdf(*_args):
        raise ModuleNotFoundError("No module named 'pypdf'", name="pypdf")

    monkeypatch.setattr(document_ingestion, "extract", missing_pypdf)
    content = pdf.read_bytes()
    response = client.post("/api/upload", files={"file": ("python_study_guide.pdf", content, "application/pdf")})
    assert response.status_code == 200, response.text
    result = response.json()
    assert "pypdf" in result["processing_warning"]
    assert result["text_extracted"] is False
    assert result["extraction_status"] == "failed"
    assert "No module named 'pypdf'" in result["extraction_error"]
    detail = client.get(f"/api/files/{result['id']}").json()
    assert detail["extraction_status"] == "failed"
    assert "pypdf" in detail["extraction_error"]
    assert client.get(f"/api/files/{result['id']}/raw", params={"dl": 1}).content == content

    monkeypatch.setattr(rag_assistant, "extract", missing_pypdf)
    answer = client.post("/api/ask", json={"question": "Summarize the uploaded guide"})
    assert answer.status_code == 200, answer.text
    assert answer.json()["answer_mode"] == "extraction_unavailable"
    assert "pypdf" in answer.json()["answer"]


def test_import_error_detail_is_visible_without_fake_extraction_success(monkeypatch):
    from services import document_ingestion

    assert client.post("/api/register", json={"email": "import-error@example.test", "password": "import-test-123"}).status_code == 200
    pdf = Path(__file__).resolve().parents[1] / "sample_documents" / "pdf" / "02_sample_project_report.pdf"

    def broken_import(*_args):
        raise ImportError("cannot import name 'PdfReader' from pypdf")

    monkeypatch.setattr(document_ingestion, "extract", broken_import)
    content = pdf.read_bytes()
    response = client.post("/api/upload", files={"file": ("import-error.pdf", content, "application/pdf")})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["text_extracted"] is False
    assert "ImportError: cannot import name 'PdfReader' from pypdf" in result["processing_warning"]
    assert client.get(f"/api/files/{result['id']}/raw", params={"dl": 1}).content == content


def test_pdf_upload_extracts_text_for_search_and_assistant():
    assert client.post("/api/register", json={"email": "pdf-extraction@example.test", "password": "pdf-test-123"}).status_code == 200
    pdf = Path(__file__).resolve().parents[1] / "sample_documents" / "pdf" / "02_sample_project_report.pdf"
    response = client.post("/api/upload", files={"file": ("python_study_guide.pdf", pdf.read_bytes(), "application/pdf")})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["text_extracted"] is True
    detail = client.get(f"/api/files/{result['id']}").json()
    assert "waterwise" in detail["text"].lower()
    search = client.get("/api/search", params={"query": "WaterWise calibration sensor"}).json()
    assert any(item["id"] == result["id"] for item in search["results"])
    answer = client.post("/api/ask", json={"question": "What does the WaterWise project measure?"}).json()
    assert any(source["id"] == result["id"] for source in answer["sources"])


def test_sample_collection_contains_readable_real_formats():
    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation
    from PIL import Image
    from pypdf import PdfReader

    sample_root = Path(__file__).resolve().parents[1] / "sample_documents"
    files = [path for path in sample_root.rglob("*") if path.is_file()]
    assert len(files) >= 14
    assert all(path.stat().st_size > 100 for path in files)

    pdfs = list((sample_root / "pdf").glob("*.pdf"))
    assert len(pdfs) >= 4
    assert all("SAMPLE" in "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
               for path in pdfs)

    docs = list((sample_root / "docx").glob("*.docx"))
    assert len(docs) >= 2
    assert all(len(Document(str(path)).paragraphs) >= 2 for path in docs)

    csv_files = list((sample_root / "csv").glob("*.csv"))
    assert len(csv_files) >= 3
    for path in csv_files:
        with path.open(encoding="utf-8", newline="") as stream:
            assert len(list(csv.reader(stream))) >= 4

    workbooks = list((sample_root / "xlsx").glob("*.xlsx"))
    assert workbooks
    workbook = load_workbook(str(workbooks[0]), read_only=True, data_only=True)
    assert {"Student Marks", "Employee Records", "Expense Data"}.issubset(set(workbook.sheetnames))
    workbook.close()

    presentations = list((sample_root / "pptx").glob("*.pptx"))
    assert presentations
    presentation = Presentation(str(presentations[0]))
    assert len(presentation.slides) >= 4
    assert any("WaterWise" in shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)

    images = list((sample_root / "images").glob("*"))
    assert {path.suffix.lower() for path in images} >= {".png", ".jpg"}
    for path in images:
        with Image.open(path) as image:
            image.verify()

    original = (sample_root / "txt" / "02_meeting_notes.txt").read_bytes()
    duplicate = (sample_root / "txt" / "03_meeting_notes_duplicate.txt").read_bytes()
    assert original == duplicate
    assert hashlib.sha256(original).digest() == hashlib.sha256(duplicate).digest()


def test_auth_document_flow_search_duplicates_comparison_and_assistant():
    email = "student@example.test"
    assert client.post("/api/register", json={"email": email, "password": "test-password-123"}).status_code == 200
    assert client.get("/api/me").json()["email"] == email

    content = b"SmartFiles test document. The project deadline is 30 November 2026."
    uploaded = client.post("/api/upload", files={"file": ("project.txt", content, "text/plain")})
    assert uploaded.status_code == 200
    file_id = uploaded.json()["id"]

    listing = client.get("/api/files")
    assert listing.status_code == 200
    listed_file = next(item for item in listing.json() if item["id"] == file_id)
    assert listed_file["extraction_status"] == "indexed"
    assert listed_file["extraction_error"] == ""
    detail = client.get(f"/api/files/{file_id}")
    assert detail.status_code == 200
    assert detail.json()["extraction_status"] == "indexed"
    assert client.get(f"/api/files/{file_id}/raw").content == content
    stats = client.get("/api/dashboard").json()
    assert stats["total"] == 1
    assert stats["extensions"][".txt"] == 1
    assert stats["categories"]
    assert client.patch(f"/api/files/{file_id}/category", json={"category": "Technical Notes"}).json()["category"] == "Technical Notes"
    assert any(item["name"] == "project.txt" and item["category"] == "Technical Notes" for item in client.get("/api/files").json())
    assert client.patch(f"/api/files/{file_id}/category", json={"category": "made up"}).status_code == 400

    found = client.get("/api/search", params={"query": "November deadline"}).json()
    assert any(item["id"] == file_id for item in found["results"])

    duplicate = client.post("/api/upload", files={"file": ("copy.txt", content, "text/plain")})
    assert duplicate.status_code == 200
    groups = client.get("/api/duplicates").json()["exact"]
    assert any(len(group["files"]) == 2 for group in groups)

    answer = client.post("/api/ask", json={"question": "What is the project deadline?"}).json()
    assert answer["sources"]
    assert answer["answer_mode"] == "extractive"
    assert answer["ai_status"] == "local_extractive"
    assert "30 November 2026" in answer["answer"]
    insufficient = client.post("/api/ask", json={"question": "What color is the spaceship?"}).json()
    assert insufficient["answer_mode"] == "insufficient"
    assert client.post("/api/ask", json={"question": "   "}).status_code == 422
    scoped = client.post("/api/ask", json={"question": "What is this document about?", "file_id": file_id}).json()
    assert scoped["answer_mode"] == "extractive"
    assert "30 November 2026" in scoped["answer"]
    assert all(source["id"] == file_id for source in scoped["sources"])

    second_content = b"SmartFiles review schedule. Final review happens in December 2026."
    second_upload = client.post("/api/upload", files={"file": ("schedule.txt", second_content, "text/plain")})
    assert second_upload.status_code == 200
    comparison = client.post("/api/compare", json={"a": file_id, "b": second_upload.json()["id"]})
    assert comparison.status_code == 200
    assert comparison.json()["diff"]
    assert comparison.json()["similarity_method"] == "TF-IDF cosine similarity"
    assert comparison.json()["a_kind"] == "TXT"
    assert 0 <= comparison.json()["similarity"] <= 1
    # When a legacy file has no saved extraction, comparison retries it from
    # the stored original and saves the recovered text for later features.
    from database.repository import q, transaction
    file_uid = q("select uid from files where id=?", (file_id,))[0]["uid"]
    with transaction() as connection:
        connection.execute("update files set text='' where id=? and uid=?", (file_id, file_uid))
    recovered = client.post("/api/compare", json={"a": file_id, "b": second_upload.json()["id"]})
    assert recovered.status_code == 200
    assert "30 November 2026" in client.get(f"/api/files/{file_id}").json()["text"]
    identical_copy = client.post("/api/upload", files={"file": ("project-copy.txt", content, "text/plain")}).json()
    identical = client.post("/api/compare", json={"a": file_id, "b": identical_copy["id"]})
    assert identical.status_code == 200
    assert identical.json()["identical"] is True
    assert identical.json()["similarity"] == 1

    assert client.post("/api/upload", files={"file": ("fake.pdf", b"not a pdf", "application/pdf")}).status_code == 400

    assert client.post("/api/register", json={"email": "other@example.test", "password": "other-password-123"}).status_code == 200
    assert client.get("/api/files").json() == []
    assert client.patch(f"/api/files/{file_id}/category", json={"category": "Other Documents"}).status_code == 404
    assert client.get(f"/api/files/{file_id}").status_code == 404
    assert client.get(f"/api/files/{file_id}/raw").status_code == 404
    assert client.post("/api/compare", json={"a": file_id, "b": second_upload.json()["id"]}).status_code == 404
    foreign_answer = client.post("/api/ask", json={"question": "What is the deadline?", "file_id": file_id})
    assert foreign_answer.status_code == 200
    assert foreign_answer.json()["sources"] == []
    foreign_metadata = client.post("/api/ask", json={"question": "Do you have project.txt?"})
    assert foreign_metadata.status_code == 200
    assert foreign_metadata.json()["answer_mode"] == "document_not_found"
    assert foreign_metadata.json()["sources"] == []

    client.post("/api/logout")
    assert client.post("/api/login", json={"email": email, "password": "test-password-123"}).status_code == 200
    assert client.delete(f"/api/files/{file_id}").status_code == 200
    assert client.get(f"/api/files/{file_id}").status_code == 404


def test_assistant_reextracts_legacy_blank_text_and_answers_from_saved_text():
    from database.repository import q, transaction

    assert client.post("/api/register", json={"email": "assistant-recovery@example.test", "password": "recovery-test-123"}).status_code == 200
    content = b"The campus shuttle departs at 8:10 AM from North Gate on weekdays. The evening shuttle leaves at 6:20 PM."
    uploaded = client.post("/api/upload", files={"file": ("shuttle_schedule.txt", content, "text/plain")})
    assert uploaded.status_code == 200, uploaded.text
    file_id = uploaded.json()["id"]
    uid = q("select uid from files where id=?", (file_id,))[0]["uid"]
    with transaction() as connection:
        connection.execute("delete from chunks where fid=? and uid=?", (file_id, uid))
        connection.execute("update files set text='',summary='' where id=? and uid=?", (file_id, uid))

    answer = client.post("/api/ask", json={"question": "When does the campus shuttle depart?"})
    assert answer.status_code == 200, answer.text
    body = answer.json()
    assert body["answer_mode"] == "extractive"
    assert "8:10 AM" in body["answer"]
    assert any(source["id"] == file_id for source in body["sources"])
    assert "8:10 AM" in client.get(f"/api/files/{file_id}").json()["text"]


def test_assistant_can_compare_two_named_documents():
    assert client.post("/api/register", json={"email": "assistant-compare@example.test", "password": "assistant-compare-123"}).status_code == 200
    first = client.post("/api/upload", files={"file": ("water_report.txt", b"Water report. Nitrate levels measured 12 milligrams per liter in spring samples.", "text/plain")})
    second = client.post("/api/upload", files={"file": ("energy_report.txt", b"Energy report. Solar panels generated 420 kilowatt hours during spring.", "text/plain")})
    assert first.status_code == second.status_code == 200
    response = client.post("/api/ask", json={"question": "Compare water_report.txt with energy_report.txt"})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["answer_mode"] == "comparison"
    assert "water_report.txt" in body["answer"] and "energy_report.txt" in body["answer"]


def test_reprocess_repairs_missing_text_and_preserves_index_on_failure(monkeypatch):
    from database.repository import q, transaction
    from services import document_ingestion

    assert client.post("/api/register", json={"email": "reprocess@example.test", "password": "reprocess-test-123"}).status_code == 200
    original = b"The Aurora lab contact email is aurora.contact@example.test. It studies urban heat islands."
    uploaded = client.post("/api/upload", files={"file": ("aurora_report.txt", original, "text/plain")})
    assert uploaded.status_code == 200
    file_id = uploaded.json()["id"]
    uid = q("select uid from files where id=?", (file_id,))[0]["uid"]
    with transaction() as connection:
        connection.execute("delete from chunks where fid=? and uid=?", (file_id, uid))
        connection.execute("update files set text='',summary='' where id=? and uid=?", (file_id, uid))

    repaired = client.post("/api/reprocess")
    assert repaired.status_code == 200, repaired.text
    result = repaired.json()
    assert result["counts"]["processed"] == 1
    assert result["processed"][0]["id"] == file_id
    repaired_detail = client.get(f"/api/files/{file_id}").json()
    assert "aurora.contact@example.test" in repaired_detail["text"]
    assert repaired_detail["extraction_status"] == "indexed"
    assert repaired_detail["extraction_error"] == ""
    assert client.post("/api/ask", json={"question": "What email address is in the Aurora report?"}).json()["answer"].find(
        "aurora.contact@example.test") >= 0

    text_before = client.get(f"/api/files/{file_id}").json()["text"]
    chunk_count_before = len(q("select id from chunks where fid=? and uid=?", (file_id, uid)))
    def fail_extraction(*_args):
        raise ImportError("deliberate test extraction failure")
    monkeypatch.setattr(document_ingestion, "extract", fail_extraction)
    failed = client.post(f"/api/reprocess?file_id={file_id}")
    assert failed.status_code == 200
    assert failed.json()["counts"]["errors"] == 1
    assert client.get(f"/api/files/{file_id}").json()["text"] == text_before
    assert len(q("select id from chunks where fid=? and uid=?", (file_id, uid))) == chunk_count_before
    assert client.post("/api/register", json={"email": "reprocess-other@example.test", "password": "other-test-123"}).status_code == 200
    assert client.post(f"/api/reprocess?file_id={file_id}").status_code == 404


def test_failed_pdf_extraction_keeps_metadata_and_reprocesses_original(monkeypatch):
    from services import document_ingestion

    assert client.post("/api/register", json={"email": "pdf-reprocess@example.test", "password": "pdf-reprocess-123"}).status_code == 200
    pdf = Path(__file__).resolve().parents[1] / "sample_documents" / "pdf" / "01_sample_resume.pdf"
    real_extract = document_ingestion.extract

    def failed_extract(*_args):
        raise ModuleNotFoundError("No module named 'pypdf'", name="pypdf")

    monkeypatch.setattr(document_ingestion, "extract", failed_extract)
    uploaded = client.post("/api/upload", files={"file": ("Vandana L Resume.pdf", pdf.read_bytes(), "application/pdf")})
    assert uploaded.status_code == 200, uploaded.text
    file_id = uploaded.json()["id"]
    assert client.get(f"/api/files/{file_id}").json()["extraction_status"] == "failed"

    metadata = client.post("/api/ask", json={"question": "Do you have Vandana's resume?"}).json()
    assert metadata["answer_mode"] == "document_found"
    assert metadata["sources"][0]["id"] == file_id

    monkeypatch.setattr(document_ingestion, "extract", real_extract)
    repair = client.post(f"/api/reprocess?file_id={file_id}")
    assert repair.status_code == 200, repair.text
    assert repair.json()["counts"]["processed"] == 1
    detail = client.get(f"/api/files/{file_id}").json()
    assert detail["extraction_status"] == "indexed"
    assert "Mira Chen" in detail["text"]


def test_frontend_chat_contract_and_local_fallback():
    assert client.post("/api/register", json={"email": "chat-contract@example.test", "password": "chat-contract-123"}).status_code == 200
    uploaded = client.post("/api/upload", files={"file": ("answer_source.txt", b"The blue comet passed Earth on 14 May 2024.", "text/plain")})
    assert uploaded.status_code == 200

    home = client.get("/")
    assistant_script = client.get("/static/js/assistant.js")
    app_script = client.get("/static/js/app.js")
    assert home.status_code == assistant_script.status_code == app_script.status_code == 200
    assert "/static/js/assistant.js" in home.text
    assert "requestJSON(API.chat" in assistant_script.text
    assert 'method: "POST"' in assistant_script.text
    assert "JSON.stringify(requestBody)" in assistant_script.text
    assert "history: assistantHistory.slice(-8)" in assistant_script.text
    assert "data.answer_mode === \"generated\"" in assistant_script.text
    assert 'credentials: "same-origin"' in app_script.text
    assert 'assistantStatus: "/api/assistant/status"' in app_script.text

    response = client.post("/api/ask", json={"question": "When did the blue comet pass Earth?"})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["answer_mode"] == "extractive"
    assert body["ai_status"] == "local_extractive"
    assert "14 May 2024" in body["answer"]
    assert body["sources"][0]["name"] == "answer_source.txt"


def test_assistant_uses_gemini_when_configured_and_falls_back_when_unavailable(monkeypatch):
    from services import llm_service
    from services import rag_assistant

    assert client.post("/api/register", json={"email": "gemini-assistant@example.test", "password": "gemini-test-123"}).status_code == 200
    uploaded = client.post("/api/upload", files={"file": ("comet.txt", b"The blue comet passed Earth on 14 May 2024.", "text/plain")})
    assert uploaded.status_code == 200

    monkeypatch.setenv("GEMINI_API_KEY", "test-gemini-key")
    monkeypatch.setattr(rag_assistant, "generate_answer", lambda question, history=None, evidence=None: "The comet date is 14 May 2024 [S1].")
    generated = client.post("/api/ask", json={"question": "When did the blue comet pass Earth?", "history": [{"role": "user", "text": "hello"}]})
    assert generated.status_code == 200, generated.text
    body = generated.json()
    assert body["answer_mode"] == "generated"
    assert body["ai_status"] == "gemini"
    assert "14 May 2024" in body["answer"]
    assert "S1" in body["citations"]

    def unavailable(*_args, **_kwargs):
        raise llm_service.ModelUnavailable("The language model timed out.")

    monkeypatch.setattr(rag_assistant, "generate_answer", unavailable)
    fallback = client.post("/api/ask", json={"question": "When did the blue comet pass Earth?"})
    assert fallback.status_code == 200, fallback.text
    fallback_body = fallback.json()
    assert fallback_body["answer_mode"] == "extractive"
    assert fallback_body["ai_status"] == "local_extractive"
    assert "14 May 2024" in fallback_body["answer"]

    status = client.get("/api/assistant/status")
    assert status.status_code == 200
    assert status.json()["configured"] is True


def test_extractor_reads_real_sample_pdf_docx_txt_csv_xlsx_and_pptx():
    from services.document_processor import extract, kind_of
    root = Path(__file__).resolve().parents[1] / "sample_documents"
    samples = [
        root / "pdf" / "02_sample_project_report.pdf",
        root / "docx" / "02_sample_project_documentation.docx",
        root / "txt" / "01_technical_notes.txt",
        root / "csv" / "student_marks.csv",
        root / "xlsx" / "fictional_records_workbook.xlsx",
        root / "pptx" / "01_waterwise_sample_presentation.pptx",
    ]
    for path in samples:
        assert path.is_file(), path
        text, _ = extract(str(path), kind_of(path.suffix.lower()), path.suffix.lower())
        assert text.strip(), path.name


def test_assistant_does_not_pick_arbitrary_document_for_unscoped_summary():
    assert client.post("/api/register", json={"email": "summary-scope@example.test", "password": "summary-test-123"}).status_code == 200
    response = client.post("/api/upload", files={"file": ("modules.txt", b"The authorization module validates sign-in sessions. The archive module stores user documents.", "text/plain")})
    assert response.status_code == 200
    body = client.post("/api/ask", json={"question": "What is this document about?"}).json()
    assert body["answer_mode"] == "insufficient"
    assert not body["sources"]


def test_supported_formats_upload_list_download_and_search():
    from io import BytesIO
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
    from pypdf import PdfWriter

    assert client.post("/api/register", json={"email": "formats@example.test", "password": "format-test-123"}).status_code == 200

    workbook = Workbook()
    workbook.active.append(["AsterismQuarterlyRevenue", "7184"])
    xlsx = BytesIO(); workbook.save(xlsx)

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    textbox.text = "BorealisPresentationLaunch 2026"
    pptx = BytesIO(); presentation.save(pptx)

    word = Document(); word.add_paragraph("CobaltDocumentSchedule 2026")
    docx = BytesIO(); word.save(docx)

    image = BytesIO(); Image.new("RGB", (8, 8), color="blue").save(image, format="PNG")
    jpeg = BytesIO(); Image.new("RGB", (8, 8), color="green").save(jpeg, format="JPEG")
    pdf = BytesIO(); writer = PdfWriter(); writer.add_blank_page(width=72, height=72); writer.write(pdf)

    files = [
        ("quarterly.csv", b"topic,value\nDandelionCsvForecast,42\n"),
        ("quarterly.xlsx", xlsx.getvalue()),
        ("launch.pptx", pptx.getvalue()),
        ("schedule.docx", docx.getvalue()),
        ("photo.png", image.getvalue()),
        ("photo.jpg", jpeg.getvalue()),
        ("photo.jpeg", jpeg.getvalue()),
        ("blank.pdf", pdf.getvalue()),
        ("notes.txt", b"EvergreenTextSearch 2026"),
    ]
    uploaded = {}
    for name, content in files:
        response = client.post("/api/upload", files={"file": (name, content, "application/octet-stream")})
        assert response.status_code == 200, (name, response.text)
        uploaded[name] = response.json()["id"]

    listing = client.get("/api/files").json()
    listed = {entry["name"]: entry for entry in listing}
    for name, content in files:
        entry = listed[name]
        assert entry["file_type"] == name.rsplit(".", 1)[1].upper()
        assert entry["size_bytes"] == len(content)
        assert entry["created"]
        assert entry["category"]
        assert entry["tags"]
        assert client.get(f"/api/files/{uploaded[name]}/raw", params={"dl": 1}).content == content

    for term, expected in [
        ("AsterismQuarterlyRevenue", "quarterly.xlsx"),
        ("BorealisPresentationLaunch", "launch.pptx"),
        ("CobaltDocumentSchedule", "schedule.docx"),
        ("DandelionCsvForecast", "quarterly.csv"),
        ("EvergreenTextSearch", "notes.txt"),
    ]:
        result = client.get("/api/search", params={"query": term}).json()
        assert any(item["name"] == expected for item in result["results"])

    # Real extracted text is compared for Office and spreadsheet uploads too.
    office_compare = client.post("/api/compare", json={"a": uploaded["quarterly.xlsx"], "b": uploaded["schedule.docx"]})
    assert office_compare.status_code == 200, office_compare.text
    assert office_compare.json()["a_kind"] == "XLSX"
    assert office_compare.json()["b_kind"] == "DOCX"
    assert client.post("/api/compare", json={"a": uploaded["blank.pdf"], "b": uploaded["notes.txt"]}).status_code == 422

    legacy = client.post("/api/upload", files={"file": ("legacy.doc", b"legacy", "application/msword")})
    assert legacy.status_code == 400 and legacy.json()["detail"].startswith("Legacy .doc")
    archive = client.post("/api/upload", files={"file": ("bundle.zip", b"PK\\x03\\x04", "application/zip")})
    assert archive.status_code == 400 and "ZIP archives are not supported" in archive.json()["detail"]
    traversal = client.post("/api/upload", files={"file": ("../safe-name.txt", b"safe", "text/plain")})
    assert traversal.status_code == 200
    assert any(entry["name"] == "safe-name.txt" for entry in client.get("/api/files").json())


def test_demo_documents_import_once_are_searchable_and_user_scoped():
    assert client.post("/api/register", json={"email": "demo@example.test", "password": "demo-test-123"}).status_code == 200

    first = client.post("/api/demo/import")
    assert first.status_code == 200, first.text
    result = first.json()
    assert result["imported"] >= 14
    assert result["errors"] == []

    auto_categories = {item["name"]: item["category"] for item in client.get("/api/files").json()}
    assert auto_categories["01_sample_resume.pdf"] == "Career & Resumes"
    assert auto_categories["01_sample_database_assignment.docx"] == "Education & Assignments"
    assert auto_categories["02_sample_project_documentation.docx"] == "Project Reports"
    assert auto_categories["02_sample_project_report.pdf"] == "Project Reports"
    assert auto_categories["03_sample_certificate.pdf"] == "Certificates"
    assert auto_categories["01_technical_notes.txt"] == "Technical Notes"
    assert auto_categories["student_marks.csv"] == "Education & Assignments"
    assert auto_categories["expense_data.csv"] == "Financial Documents"
    assert auto_categories["01_waterwise_sample_presentation.pptx"] == "Presentations"
    assert auto_categories["waterwise_data_flow.png"] == "Images & Scanned Documents"

    listed = client.get("/api/files").json()
    assert len(listed) == result["imported"]
    assert all(item["is_demo"] for item in listed)
    assert client.get("/api/dashboard").json()["demo_total"] == result["imported"]

    again = client.post("/api/demo/import")
    assert again.status_code == 200
    assert again.json()["imported"] == 0
    assert again.json()["skipped"] == result["imported"]

    search = client.get("/api/search", params={"query": "soil moisture below 28 percent"}).json()
    assert any(item["is_demo"] and "technical_notes" in item["name"] for item in search["results"])
    duplicates = client.get("/api/duplicates").json()["exact"]
    duplicate_names = [item["name"] for group in duplicates for item in group["files"]]
    assert "02_meeting_notes.txt" in duplicate_names
    assert "03_meeting_notes_duplicate.txt" in duplicate_names
    assert all(item["is_demo"] for group in duplicates for item in group["files"])
    answer = client.post("/api/ask", json={"question": "What does the dashboard flag below 28 percent?"}).json()
    assert answer["sources"] and any(source["is_demo"] for source in answer["sources"])
    notes = client.post("/api/ask", json={"question": "What are the main points in my notes?"}).json()
    assert notes["answer_mode"] == "extractive"
    assert any("notes" in source["name"].lower() for source in notes["sources"])

    assert client.post("/api/register", json={"email": "another-demo@example.test", "password": "demo-test-456"}).status_code == 200
    assert client.get("/api/files").json() == []
    assert client.get("/api/dashboard").json()["demo_total"] == 0


def test_protected_routes_require_login():
    client.post("/api/logout")
    assert client.get("/api/files").status_code == 401
    assert client.patch("/api/files/1/category", json={"category": "Other Documents"}).status_code == 401
    assert client.get("/api/duplicates").status_code == 401
    assert client.post("/api/ask", json={"question": "Anything?"}).status_code == 401
    assert client.get("/api/compression/history").status_code == 401
    assert client.get("/api/storage-optimizer").status_code == 401
    assert client.get("/api/me").status_code == 401
    assert client.get("/docs").status_code == 200


def test_assistant_answers_education_question_from_resume_and_greeting_is_clean(monkeypatch):
    from io import BytesIO
    from PIL import Image
    from services import rag_assistant

    assert client.post("/api/register", json={"email": "education-answer@example.test", "password": "education-test-123"}).status_code == 200
    resume = ("RESUME\nAsha Rao, MCA student\nSkills: Python, FastAPI, SQL.\n"
              "Education: Master of Computer Applications, Visvesvaraya Technological University, 2025 to 2027.\n"
              "Experience: Python web development intern at BrightLeaf Technologies.")
    uploaded = client.post("/api/upload", files={"file": ("asha_rao_resume.txt", resume.encode(), "text/plain")})
    assert uploaded.status_code == 200, uploaded.text

    png = BytesIO()
    Image.new("RGB", (8, 8), color="white").save(png, format="PNG")
    image_upload = client.post("/api/upload", files={"file": ("unrelated_certificate.png", png.getvalue(), "image/png")})
    assert image_upload.status_code == 200, image_upload.text

    def fail_if_ocr_is_attempted(*_args):
        raise AssertionError("A greeting or unrelated resume question must not retry OCR on this image")

    monkeypatch.setattr(rag_assistant, "extract", fail_if_ocr_is_attempted)
    greeting = client.post("/api/ask", json={"question": "hello"})
    assert greeting.status_code == 200
    assert greeting.json()["answer_mode"] == "greeting"
    assert greeting.json()["sources"] == []
    assert greeting.json()["extraction_warnings"] == []

    response = client.post("/api/ask", json={"question": "where did asha roa study her mca"})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["answer_mode"] == "extractive"
    assert "Visvesvaraya Technological University" in body["answer"]
    assert [source["name"] for source in body["sources"]] == ["asha_rao_resume.txt"]
    assert body["extraction_warnings"] == []


def test_assistant_finds_documents_from_metadata_before_content_search():
    assert client.post("/api/register", json={"email": "metadata-resume@example.test", "password": "metadata-test-123"}).status_code == 200
    resume_path = Path(__file__).resolve().parents[1] / "sample_documents" / "pdf" / "01_sample_resume.pdf"
    content = resume_path.read_bytes()
    uploaded = client.post("/api/upload", files={"file": ("Vandana L Resume.pdf", content, "application/pdf")})
    assert uploaded.status_code == 200, uploaded.text
    document_id = uploaded.json()["id"]
    second_resume = ("RESUME\nLeena Park is the fictional candidate for this resume document.\nEducation: Master of Science in analytics at North Coast University.\n"
                     "Experience: Leena built data quality dashboards for a fictional research group.")
    second_upload = client.post("/api/upload", files={"file": ("leena_resume.txt", second_resume.encode(), "text/plain")})
    assert second_upload.status_code == 200, second_upload.text
    second_id = second_upload.json()["id"]

    for question in (
        "Do you have Vandana's resume?",
        "Do you have Vandna's CV?",
        "Do I have a PDF?",
        "Find my resume.",
        "What is the name of my resume file?",
    ):
        result = client.post("/api/ask", json={"question": question})
        assert result.status_code == 200, result.text
        body = result.json()
        assert body["answer_mode"] == "document_found"
        assert "Vandana L Resume.pdf" in body["answer"]
        assert any(source["name"] == "Vandana L Resume.pdf" and source["file_type"] == "PDF" for source in body["sources"])
        assert document_id in [source["id"] for source in body["sources"]]
        assert all(client.get(f"/api/files/{source['id']}").status_code == 200 for source in body["sources"])

    summary = client.post("/api/ask", json={"question": "Summarize my resume."})
    assert summary.status_code == 200, summary.text
    summary_body = summary.json()
    assert summary_body["answer_mode"] == "extractive"
    assert summary_body["sources"]
    assert {source["id"] for source in summary_body["sources"]} == {document_id, second_id}
    assert "Mira Chen" in summary_body["answer"]
    assert "Leena Park" in summary_body["answer"]

    absent = client.post("/api/ask", json={"question": "Do I have a document about a topic that is not present?"})
    assert absent.status_code == 200, absent.text
    absent_body = absent.json()
    assert absent_body["answer_mode"] == "document_not_found"
    assert "couldn't find a matching document" in absent_body["answer"].lower()
    assert absent_body["sources"] == []


def test_search_supports_metadata_and_upload_date_filters():
    assert client.post("/api/register", json={"email": "search-filters@example.test", "password": "search-test-123"}).status_code == 200
    payload = b"Network security and encrypted storage notes for the project team."
    uploaded = client.post("/api/upload", files={"file": ("security_notes.txt", payload, "text/plain")})
    assert uploaded.status_code == 200, uploaded.text
    fid = uploaded.json()["id"]
    assert client.patch(f"/api/files/{fid}/category", json={"category": "Technical Notes"}).status_code == 200
    params = {"category": "technical notes", "tag": "security", "file_type": "TXT"}
    result = client.get("/api/search", params=params)
    assert result.status_code == 200, result.text
    assert [item["id"] for item in result.json()["results"]] == [fid]
    assert client.get("/api/search", params={**params, "file_type": "pdf"}).json()["results"] == []
    assert client.get("/api/search", params={**params, "date_from": 9999999999}).json()["results"] == []
    assert client.get("/api/search", params={"date_from": 200, "date_to": 100}).status_code == 422


def test_near_duplicate_text_detection_returns_explainable_pair():
    assert client.post("/api/register", json={"email": "near-duplicates@example.test", "password": "near-test-123"}).status_code == 200
    base = ("The fictional research group prepared a document management project using Python FastAPI SQLite secure storage "
            "and searchable records for university students and staff. The application tracks uploads, categories, tags, "
            "summaries, and access permissions for each account. " * 5)
    changed = base.replace("university students and staff", "university students and faculty members") + "The team also reviewed backup procedures."
    for name, content in (("research_report_one.txt", base), ("research_report_two.txt", changed)):
        response = client.post("/api/upload", files={"file": (name, content.encode(), "text/plain")})
        assert response.status_code == 200, response.text
    result = client.get("/api/duplicates")
    assert result.status_code == 200, result.text
    pairs = result.json()["near_duplicates"]
    assert len(pairs) == 1
    assert pairs[0]["score"] >= .72
    assert "Jaccard" in pairs[0]["metric"]
    assert {f["name"] for f in pairs[0]["files"]} == {"research_report_one.txt", "research_report_two.txt"}


def test_dashboard_reports_processing_and_duplicate_counts():
    assert client.post("/api/register", json={"email": "dashboard-status@example.test", "password": "dashboard-test-123"}).status_code == 200
    content = b"This actual readable note contains enough technical details about secure database indexing and retrieval."
    first = client.post("/api/upload", files={"file": ("status_notes.txt", content, "text/plain")})
    second = client.post("/api/upload", files={"file": ("status_notes_copy.txt", content, "text/plain")})
    assert first.status_code == second.status_code == 200
    body = client.get("/api/dashboard").json()
    assert body["total"] == 2
    assert body["processing"]["indexed"] == 2
    assert body["duplicates"]["exact_groups"] == 1
    assert body["duplicates"]["potential_near_pairs"] == 0


def test_repeated_upload_reuses_successful_extraction(monkeypatch):
    assert client.post("/api/register", json={"email": "reuse-text@example.test", "password": "reuse-test-123"}).status_code == 200
    from services import document_ingestion
    content = b"This fictional document stores readable text about software testing and document indexing for this account."
    first = client.post("/api/upload", files={"file": ("reuse_notes.txt", content, "text/plain")})
    assert first.status_code == 200
    def unexpected_extraction(*_args):
        raise AssertionError("unchanged content should use the cached extracted text")
    monkeypatch.setattr(document_ingestion, "extract", unexpected_extraction)
    second = client.post("/api/upload", files={"file": ("reuse_notes_copy.txt", content, "text/plain")})
    assert second.status_code == 200, second.text
    assert second.json()["extraction_reused"] is True
    assert second.json()["text_extracted"] is True


def test_long_document_summary_uses_real_sections():
    from services.summarizer import summarize
    first = "The fictional report describes a marine biology research survey at North Bay University. "
    second = "The project recorded coral temperature readings from August through October and stored results in a database. "
    summary = summarize((first * 350) + (second * 350))
    assert "marine biology" in summary.lower()
    assert "temperature readings" in summary.lower()


def test_pdf_compression_preserves_original_and_reports_real_sizes():
    from io import BytesIO
    from pypdf import PdfReader

    assert client.post("/api/register", json={"email": "compress-pdf@example.test", "password": "compress-test-123"}).status_code == 200
    sample = Path(__file__).resolve().parents[1] / "sample_documents" / "pdf" / "02_sample_project_report.pdf"
    uploaded = client.post("/api/upload", files={"file": (sample.name, sample.read_bytes(), "application/pdf")})
    assert uploaded.status_code == 200, uploaded.text
    file_id = uploaded.json()["id"]
    original = client.get(f"/api/files/{file_id}/raw").content

    response = client.post("/api/compress", data={"file_id": str(file_id)})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["original_size_bytes"] == len(original)
    assert result["compressed_size_bytes"] > 0
    assert result["space_saved_bytes"] == len(original) - result["compressed_size_bytes"]
    expected_percent = (result["space_saved_bytes"] / len(original)) * 100
    assert result["percent_reduction"] == round(expected_percent, 2)
    compressed = client.get(result["download_url"])
    assert compressed.status_code == 200
    assert len(PdfReader(BytesIO(compressed.content)).pages) > 0
    assert client.get(f"/api/files/{file_id}/raw").content == original
    assert client.get("/api/compression/history").json()[0]["id"] == result["id"]

    dashboard = client.get("/api/dashboard").json()
    assert dashboard["storage"]["original_bytes"] == len(original)
    assert dashboard["storage"]["compressed_bytes"] == result["compressed_size_bytes"]
    assert dashboard["storage"]["confirmed_saved_bytes"] == max(0, result["space_saved_bytes"])
    assert sum(item["count"] for item in dashboard["monthly_uploads"]) == 1
    filtered = client.get("/api/dashboard", params={"file_type": "pdf"}).json()
    assert filtered["total"] == 1
    assert filtered["extensions"] == {".pdf": 1}
    assert client.get("/api/dashboard", params={"file_type": "png"}).json()["total"] == 0
    category = client.get("/api/files").json()[0]["category"]
    assert client.get("/api/dashboard", params={"category": category}).json()["total"] == 1
    assert client.get("/api/dashboard", params={"date_from": "9999999999"}).json()["total"] == 0
    optimizer = client.get("/api/storage-optimizer").json()
    assert optimizer["compressible_documents"][0]["id"] == file_id
    assert optimizer["estimated_savings_bytes"] is None


def test_image_compression_stores_valid_private_result_and_handles_no_savings(monkeypatch):
    from io import BytesIO
    from PIL import Image
    from fastapi.testclient import TestClient
    from app import app
    from services import file_compression

    assert client.post("/api/register", json={"email": "compress-image@example.test", "password": "image-test-123"}).status_code == 200
    image = Image.new("RGB", (512, 512), (17, 82, 139))
    original_buffer = BytesIO()
    image.save(original_buffer, format="PNG", compress_level=0)
    original = original_buffer.getvalue()
    response = client.post(
        "/api/compress",
        files={"file": ("diagram.png", original, "image/png")},
        data={"quality": "85", "max_dimension": "0"},
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["status"] == "compressed"
    assert result["compressed_size_bytes"] < len(original)
    preview = client.get(result["preview_url"])
    assert preview.status_code == 200
    with Image.open(BytesIO(preview.content)) as compressed_image:
        assert compressed_image.size == (512, 512)

    resized = client.post(
        "/api/compress",
        files={"file": ("resized.png", original, "image/png")},
        data={"max_dimension": "256"},
    )
    assert resized.status_code == 200, resized.text
    with Image.open(BytesIO(client.get(resized.json()["preview_url"]).content)) as resized_image:
        assert resized_image.size == (256, 256)

    other_client = TestClient(app)
    assert other_client.post("/api/register", json={"email": "compress-other@example.test", "password": "other-test-123"}).status_code == 200
    assert other_client.get(result["download_url"]).status_code == 404
    assert other_client.get("/api/compression/history").json() == []

    tiny = BytesIO()
    Image.new("RGB", (2, 2), "white").save(tiny, format="JPEG", quality=20)
    larger_output = BytesIO()
    Image.new("RGB", (2, 2), "white").save(larger_output, format="JPEG", quality=95)
    monkeypatch.setattr(file_compression, "_compress_image", lambda *_args: larger_output.getvalue() + (b"\x00" * 2048))
    no_savings = client.post("/api/compress", files={"file": ("tiny.jpg", tiny.getvalue(), "image/jpeg")})
    assert no_savings.status_code == 200, no_savings.text
    body = no_savings.json()
    assert body["status"] == "not_reduced"
    assert body["compressed_size_bytes"] == body["original_size_bytes"]
    assert body["space_saved_bytes"] == 0
    assert body["percent_reduction"] == 0
    assert "original is available unchanged" in body["message"]
    assert client.get(body["download_url"]).content == tiny.getvalue()


def test_compression_rejects_unsupported_corrupted_and_password_protected_files():
    from io import BytesIO
    from pypdf import PdfWriter

    assert client.post("/api/register", json={"email": "compress-errors@example.test", "password": "errors-test-123"}).status_code == 200
    unsupported = client.post("/api/compress", files={"file": ("archive.zip", b"PK", "application/zip")})
    assert unsupported.status_code == 415
    corrupted = client.post("/api/compress", files={"file": ("broken.pdf", b"%PDF-1.7 not a real pdf", "application/pdf")})
    assert corrupted.status_code == 422
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("secret")
    encrypted = BytesIO()
    writer.write(encrypted)
    password_protected = client.post("/api/compress", files={"file": ("locked.pdf", encrypted.getvalue(), "application/pdf")})
    assert password_protected.status_code == 422
    assert "Password-protected" in password_protected.json()["detail"]
    assert client.get("/api/compression/history").json() == []


def test_compression_modes_cover_text_scanned_pdf_large_jpeg_and_alpha_png():
    from io import BytesIO
    import fitz
    from PIL import Image

    assert client.post("/api/register", json={"email": "compress-modes@example.test", "password": "modes-test-123"}).status_code == 200

    # Digitally generated PDF: text remains selectable after rewriting.
    text_doc = fitz.open()
    page = text_doc.new_page()
    page.insert_text((72, 72), "SmartFiles keeps searchable PDF text intact.", fontsize=16)
    text_pdf = text_doc.tobytes()
    text_doc.close()
    text_result = client.post("/api/compress", files={"file": ("text-report.pdf", text_pdf, "application/pdf")}, data={"mode": "lossless"})
    assert text_result.status_code == 200, text_result.text
    text_download = client.get(text_result.json()["download_url"]).content
    with fitz.open(stream=text_download, filetype="pdf") as checked:
        assert "searchable PDF text" in checked[0].get_text()
    assert text_result.json()["original_size_bytes"] == len(text_pdf)
    assert text_result.json()["space_saved_bytes"] == max(0, len(text_pdf) - len(text_download))

    # Scanned/image-heavy PDF: downsample the embedded scan without rasterizing
    # the page or replacing its PDF structure.
    scan = Image.effect_noise((2600, 2000), 75).convert("RGB")
    scan_bytes = BytesIO()
    scan.save(scan_bytes, "JPEG", quality=96)
    image_pdf_doc = fitz.open()
    image_page = image_pdf_doc.new_page(width=612, height=792)
    image_page.insert_image(image_page.rect, stream=scan_bytes.getvalue())
    image_pdf = image_pdf_doc.tobytes()
    image_pdf_doc.close()
    scan_result = client.post("/api/compress", files={"file": ("scanned-report.pdf", image_pdf, "application/pdf")}, data={"mode": "balanced"})
    assert scan_result.status_code == 200, scan_result.text
    scan_download = client.get(scan_result.json()["download_url"]).content
    with fitz.open(stream=scan_download, filetype="pdf") as checked:
        assert checked.page_count == 1
        assert checked[0].get_images()
    assert len(scan_download) < len(image_pdf)

    # Large JPEG and transparent PNG both return decodable images, and the PNG
    # retains its alpha channel through palette optimization.
    large_jpeg = BytesIO()
    Image.effect_noise((2400, 1800), 65).convert("RGB").save(large_jpeg, "JPEG", quality=96)
    jpeg_result = client.post("/api/compress", files={"file": ("photo.jpg", large_jpeg.getvalue(), "image/jpeg")}, data={"mode": "balanced"})
    assert jpeg_result.status_code == 200, jpeg_result.text
    jpeg_download = client.get(jpeg_result.json()["download_url"]).content
    assert len(jpeg_download) < len(large_jpeg.getvalue())
    with Image.open(BytesIO(jpeg_download)) as checked:
        checked.verify()

    png = Image.new("RGBA", (1200, 900))
    pixels = png.load()
    for y in range(png.height):
        for x in range(png.width):
            pixels[x, y] = (x % 256, y % 256, (x + y) % 256, (x * y) % 256)
    png_bytes = BytesIO()
    png.save(png_bytes, "PNG", compress_level=0)
    png_result = client.post("/api/compress", files={"file": ("transparent.png", png_bytes.getvalue(), "image/png")}, data={"mode": "balanced"})
    assert png_result.status_code == 200, png_result.text
    png_download = client.get(png_result.json()["download_url"]).content
    with Image.open(BytesIO(png_download)) as checked:
        checked.load()
        assert checked.convert("RGBA").getchannel("A").getextrema() == (0, 255)

    for result, original, downloaded in ((scan_result.json(), image_pdf, scan_download),
                                         (jpeg_result.json(), large_jpeg.getvalue(), jpeg_download),
                                         (png_result.json(), png_bytes.getvalue(), png_download)):
        assert result["original_size_bytes"] == len(original)
        assert result["compressed_size_bytes"] == len(downloaded)
        assert result["space_saved_bytes"] == len(original) - len(downloaded) > 0
        assert result["percent_reduction"] == round((len(original) - len(downloaded)) / len(original) * 100, 2)


def test_dashboard_and_compression_views_are_connected_to_frontend_assets():
    home = client.get("/")
    assert home.status_code == 200
    assert '/static/favicon.svg' in home.text
    assert client.get("/static/favicon.svg").status_code == 200
    for view in ("comparisonView", "plagiarismView", "compressorView", "optimizerView", "searchView"):
        assert f'id="{view}"' in home.text
    assert 'data-view="compressor"' in home.text
    assert 'data-view="optimizer"' in home.text
    assert 'data-view="plagiarism"' in home.text
    assert '<input id="plagiarismFile" type="file" accept=".pdf,.docx,.txt">' in home.text
    assert '<details class="plagiarism-paste-option">' in home.text
    assert 'id="plagiarismText"' in home.text
    assert 'id="plagiarismScanButton"' in home.text
    assert "Check Plagiarism" in home.text
    assert 'id="plagiarismReport"' in home.text
    for removed_ui in (
        'id="plagiarismHistory"', 'id="plagiarismComparison"',
        'id="comparePlagiarismReports"', 'id="loadMorePlagiarismReports"',
        "Previous reports", "View Report", "plagiarismHistorySearch",
    ):
        assert removed_ui not in home.text
    for asset in ("analytics.js", "compressor.js", "plagiarism.js"):
        response = client.get(f"/static/js/{asset}")
        assert response.status_code == 200
        assert "API." in response.text
    plagiarism_script = client.get("/static/js/plagiarism.js").text
    assert 'online_search: true' in plagiarism_script
    assert 'source.verification === "provider_match"' in plagiarism_script
    assert 'source.evidence' in plagiarism_script
    assert 'match.title' in plagiarism_script
    assert 'link.href = match.url' in plagiarism_script
    assert "No matching content was found by Quetext." in plagiarism_script
    assert "Quetext could not complete the online check." in plagiarism_script
    assert "Free local comparison completed against your saved SmartFiles documents." in plagiarism_script
    stylesheet = client.get("/static/css/style.css").text
    assert ".plagiarism-match-card" in stylesheet
    assert '<script src="https://cdn.' not in home.text


def test_plagiarism_history_keeps_missing_metrics_unavailable(monkeypatch):
    from services import plagiarism

    monkeypatch.setattr(plagiarism, "q", lambda *_args: [{
        "id": "legacy-scan",
        "fid": 17,
        "document_name": "legacy.txt",
        "status": "completed",
        "created": "2026-01-01T00:00:00",
        "updated": "2026-01-01T00:00:00",
        "report_json": '{"scan_date":"2026-01-01T00:00:00"}',
    }])

    report = plagiarism.list_reports("owner")[0]

    assert report["estimated_text_similarity_percent"] is None
    assert report["approximate_text_similarity_percent"] is None
    assert report["exact_match_count"] is None
    assert report["approximate_match_count"] is None


def test_plagiarism_history_is_paginated_and_reports_saved_coverage():
    assert client.post("/api/register", json={"email": "plagiarism-history-page@example.test", "password": "history-page-123"}).status_code == 200
    uploaded = client.post(
        "/api/upload",
        files={"file": ("history_page.txt", b"An analyzed passage contains enough words for a reliable local comparison.", "text/plain")},
    )
    file_id = uploaded.json()["id"]
    scan_ids = []
    for _ in range(3):
        started = client.post("/api/plagiarism/scans", json={"file_id": file_id})
        assert started.status_code == 202
        scan_ids.append(started.json()["scan_id"])

    first_page = client.get("/api/plagiarism/reports", params={"limit": 2, "offset": 0})
    second_page = client.get("/api/plagiarism/reports", params={"limit": 2, "offset": 2})
    assert first_page.status_code == 200
    assert len(first_page.json()) == 2
    assert len(second_page.json()) == 1
    assert len({item["scan_id"] for item in first_page.json() + second_page.json()}) == 3
    latest = first_page.json()[0]
    assert latest["file_id"] == file_id
    assert latest["eligible_words_analyzed"] > 0
    assert latest["analyzed_passages"] == 1
    assert latest["skipped_passages"] == 0
    assert latest["exact_match_count"] == 0
    assert client.get("/api/plagiarism/reports", params={"limit": 101}).status_code == 422
    local_comparison = client.get(
        f"/api/plagiarism/comparisons/{scan_ids[0]}/{scan_ids[1]}"
    ).json()
    assert local_comparison["online_matches"] is None
    assert local_comparison["sources"] is None


def test_plagiarism_scan_comparison_tracks_verified_changes_and_warns_on_coverage():
    import json
    from database.connection import DB

    assert client.post("/api/register", json={"email": "plagiarism-compare@example.test", "password": "compare-test-123"}).status_code == 200
    uploaded = client.post(
        "/api/upload",
        files={"file": ("compare_draft.txt", b"Distinctive wording in this passage provides a consistent comparison target.", "text/plain")},
    )
    file_id = uploaded.json()["id"]
    earlier_id = client.post("/api/plagiarism/scans", json={"file_id": file_id}).json()["scan_id"]
    later_id = client.post("/api/plagiarism/scans", json={"file_id": file_id}).json()["scan_id"]
    earlier = client.get(f"/api/plagiarism/reports/{earlier_id}").json()
    later = client.get(f"/api/plagiarism/reports/{later_id}").json()

    earlier.update({
        "online_search_requested": True,
        "online_search_status": "completed",
        "searches_completed": 1,
        "online_pages_analyzed": 1,
        "skipped_online_passages": 2,
        "sources": [{
            "title": "First source", "url": "https://example.org/first",
            "verification": "retrieved_page_text", "matches": [],
        }],
        "online_matches": [{
            "passage": "Distinctive wording in this passage provides a consistent comparison target.",
            "match_type": "exact",
            "sources": [{
                "title": "First source", "url": "https://example.org/first",
                "evidence": "Distinctive wording in this passage provides a consistent comparison target.",
            }],
        }],
    })
    later.update({
        "online_search_requested": True,
        "online_search_status": "completed",
        "searches_completed": 1,
        "online_pages_analyzed": 2,
        "skipped_online_passages": 1,
        "local_matches": [{
            "document_id": 987, "document_name": "reference.txt", "matches": [{
                "passage": "Distinctive wording in this passage provides a consistent comparison target.",
                "match_type": "exact",
                "evidence": "Distinctive wording in this passage provides a consistent comparison target.",
            }],
        }],
        "sources": [{
            "title": "Second source", "url": "https://example.org/second",
            "verification": "retrieved_page_text", "matches": [],
        }],
        "online_matches": [{
            "passage": "Distinctive wording in this passage provides a consistent comparison target.",
            "match_type": "exact",
            "sources": [{
                "title": "Second source", "url": "https://example.org/second",
                "evidence": "Distinctive wording in this passage provides a consistent comparison target.",
            }],
        }],
    })
    with DB:
        DB.execute("UPDATE plagiarism_scans SET report_json=? WHERE id=?", (json.dumps(earlier), earlier_id))
        DB.execute("UPDATE plagiarism_scans SET report_json=? WHERE id=?", (json.dumps(later), later_id))

    comparison = client.get(f"/api/plagiarism/comparisons/{earlier_id}/{later_id}")
    assert comparison.status_code == 200, comparison.text
    result = comparison.json()
    assert result["metrics"]["online_pages_checked"] == {"earlier": 1, "later": 2, "change": 1}
    assert result["metrics"]["online_unchecked_passages"] == {"earlier": 2, "later": 1, "change": -1}
    assert result["local_matches"]["newly_detected"][0]["source_name"] == "reference.txt"
    assert result["online_matches"]["newly_detected"][0]["source_name"] == "Second source"
    assert result["sources"]["newly_returned"][0]["url"] == "https://example.org/second"
    assert any("unchecked passage counts differ" in warning for warning in result["warnings"])
    assert any("online pages with returned text differs" in warning for warning in result["warnings"])
    assert "do not mean the document became more original" in result["interpretation"]

    other_document = client.post(
        "/api/upload",
        files={"file": ("different_document.txt", b"This is a different uploaded document for comparison.", "text/plain")},
    )
    different_scan_id = client.post("/api/plagiarism/scans", json={"file_id": other_document.json()["id"]}).json()["scan_id"]
    assert client.get(f"/api/plagiarism/comparisons/{earlier_id}/{different_scan_id}").status_code == 422
    assert client.get(f"/api/plagiarism/comparisons/{earlier_id}/{earlier_id}").status_code == 422


def test_plagiarism_local_exact_approximate_and_report_persistence(monkeypatch):
    from database.connection import DB
    from services import plagiarism

    monkeypatch.setenv("QUETEXT_API_KEY", "must-not-be-used-for-local-scans")
    monkeypatch.setattr(plagiarism, "_quetext_request", lambda *_args, **_kwargs: pytest.fail("A local-only scan must not call an online API"))
    assert client.post("/api/register", json={"email": "plagiarism-local@example.test", "password": "plagiarism-test-123"}).status_code == 200
    exact = "SmartFiles compares each document passage with private workspace files to find repeated wording."
    approximate_target = "Carefully compare matching passages because text overlap alone cannot prove that authors copied content."
    approximate_source = "Carefully compare matching passages because word overlap alone cannot prove that authors copied content."
    reference = client.post(
        "/api/upload",
        files={"file": ("reference_notes.txt", f"{exact}\n{approximate_source}".encode(), "text/plain")},
    )
    draft = client.post(
        "/api/upload",
        files={"file": ("draft_report.txt", f"{exact}\n{approximate_target}".encode(), "text/plain")},
    )
    assert reference.status_code == 200
    assert draft.status_code == 200

    started = client.post("/api/plagiarism/scans", json={"file_id": draft.json()["id"]})
    assert started.status_code == 202, started.text
    report_response = client.get(f"/api/plagiarism/reports/{started.json()['scan_id']}")
    assert report_response.status_code == 200, report_response.text
    report = report_response.json()
    assert report["status"] == "completed"
    assert report["exact_match_count"] == 1
    assert report["approximate_match_count"] == 1
    analyzed_words = len(plagiarism._words(exact)) + len(plagiarism._words(approximate_target))
    assert report["estimated_text_similarity_percent"] == round(100 * len(plagiarism._words(exact)) / analyzed_words, 1)
    assert report["approximate_text_similarity_percent"] == round(100 * len(plagiarism._words(approximate_target)) / analyzed_words, 1)
    assert report["unmatched_text_percent"] < 100
    assert report["local_matches"][0]["document_name"] == "reference_notes.txt"
    exact_detail = next(match for match in report["matches"] if match["exact_sources"])
    approximate_detail = next(match for match in report["matches"] if match["approximate_sources"])
    assert exact_detail["passage"] == exact
    assert exact_detail["exact_sources"][0]["evidence"] == exact
    assert exact_detail["matched_text"] == exact
    assert exact_detail["exact_sources"][0]["type"] == "local"
    assert exact_detail["exact_sources"][0]["document_id"] == reference.json()["id"]
    assert exact_detail["exact_sources"][0]["source_start_offset"] == 0
    assert approximate_detail["passage"] == approximate_target
    assert approximate_detail["approximate_sources"][0]["evidence"] == approximate_source
    assert 0.78 <= approximate_detail["approximate_sources"][0]["word_jaccard"] <= 1
    assert approximate_detail["approximate_sources"][0]["source_type"] == "local"
    assert approximate_detail["approximate_sources"][0]["source_text"] == approximate_source
    assert "Approximate-only coverage" in report["method"]
    assert "Text similarity does not by itself prove plagiarism" in report["disclaimer"]
    assert client.get("/api/plagiarism/reports").json()[0]["scan_id"] == started.json()["scan_id"]

    persisted = DB.execute(
        "SELECT report_json FROM plagiarism_scans WHERE id=?",
        (started.json()["scan_id"],),
    ).fetchone()
    assert persisted is not None
    assert '"document_name": "draft_report.txt"' in persisted["report_json"]


def test_plagiarism_candidate_passage_sampling_preserves_dictionary_records(monkeypatch):
    from services import plagiarism

    monkeypatch.setattr(plagiarism, "MAX_ANALYZED_PASSAGES", 3)
    passages, total = plagiarism._candidate_passages([
        "Alpha beta gamma delta epsilon zeta.",
        "Bravo charlie delta echo foxtrot golf.",
        "Hotel india juliet kilo lima mike.",
        "November oscar papa quebec romeo sierra.",
        "Tango uniform victor whiskey xray yankee.",
    ])

    assert total == 5
    assert len(passages) == 3
    assert all(isinstance(passage, dict) for passage in passages)
    assert len({passage["index"] for passage in passages}) == len(passages)
    assert [passage["text"] for passage in passages] == [
        "Alpha beta gamma delta epsilon zeta.",
        "Bravo charlie delta echo foxtrot golf.",
        "Hotel india juliet kilo lima mike.",
    ]


def test_plagiarism_long_document_scan_completes_after_sampling(monkeypatch):
    from services import plagiarism

    monkeypatch.setattr(plagiarism, "MAX_ANALYZED_PASSAGES", 3)
    assert client.post("/api/register", json={
        "email": "plagiarism-long-scan@example.test",
        "password": "long-scan-test-123",
    }).status_code == 200
    text = "\n".join([
        "Alpha beta gamma delta epsilon zeta.",
        "Bravo charlie delta echo foxtrot golf.",
        "Hotel india juliet kilo lima mike.",
        "November oscar papa quebec romeo sierra.",
        "Tango uniform victor whiskey xray yankee.",
    ])
    uploaded = client.post(
        "/api/upload",
        files={"file": ("long_scan.txt", text.encode(), "text/plain")},
    )
    assert uploaded.status_code == 200

    started = client.post("/api/plagiarism/scans", json={"file_id": uploaded.json()["id"]})
    report = client.get(f"/api/plagiarism/reports/{started.json()['scan_id']}").json()

    assert report["status"] == "completed_with_warnings"
    assert report["eligible_words_analyzed"] == 18
    assert report["analyzed_passages"] == 3


def test_plagiarism_scan_failure_does_not_expose_internal_error(monkeypatch):
    from services import plagiarism

    assert client.post("/api/register", json={
        "email": "plagiarism-scan-error@example.test",
        "password": "scan-error-test-123",
    }).status_code == 200
    uploaded = client.post(
        "/api/upload",
        files={"file": ("scan_error.txt", b"A readable passage contains enough words for plagiarism analysis.", "text/plain")},
    )
    assert uploaded.status_code == 200

    def fail_analysis(*_args):
        raise TypeError("internal dictionary failure detail")

    monkeypatch.setattr(plagiarism, "_candidate_passages", fail_analysis)
    started = client.post("/api/plagiarism/scans", json={"file_id": uploaded.json()["id"]})
    report = client.get(f"/api/plagiarism/reports/{started.json()['scan_id']}").json()

    assert report["status"] == "failed"
    assert any("internal processing error" in warning for warning in report["errors"])
    assert "TypeError" not in str(report)
    assert "internal dictionary failure detail" not in str(report)


def test_online_plagiarism_pdf_excludes_other_document_matches():
    from services.plagiarism_pdf import _local_matches

    report = {
        "online_search_requested": True,
        "local_matches": [{"document_name": "another.pdf", "matches": [{
            "passage": "A passage from a different PDF.", "evidence": "The other PDF text.",
        }]}],
    }
    assert list(_local_matches(report)) == []


def test_plagiarism_pdf_download_uses_saved_report_and_flows_long_passages(monkeypatch):
    from io import BytesIO
    from pypdf import PdfReader

    assert client.post("/api/register", json={"email": "plagiarism-pdf@example.test", "password": "pdf-test-123"}).status_code == 200
    passage = "Password=hunter2 " + " ".join(f"passageword{index}" for index in range(1400)) + "."
    reference = client.post(
        "/api/upload",
        files={"file": ("pdf_reference.txt", passage.encode(), "text/plain")},
    )
    draft = client.post(
        "/api/upload",
        files={"file": ("pdf_draft.txt", passage.encode(), "text/plain")},
    )
    assert reference.status_code == 200
    assert draft.status_code == 200
    started = client.post("/api/plagiarism/scans", json={"file_id": draft.json()["id"]})
    scan_id = started.json()["scan_id"]

    downloaded = client.get(f"/api/plagiarism/reports/{scan_id}/pdf")
    assert downloaded.status_code == 200, downloaded.text
    assert downloaded.headers["content-type"] == "application/pdf"
    assert downloaded.headers["content-disposition"] == f'attachment; filename="SmartFiles_plagiarism_report_{scan_id}.pdf"'
    assert downloaded.headers["cache-control"] == "no-store"
    pdf = PdfReader(BytesIO(downloaded.content))
    assert len(pdf.pages) > 1
    text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "pdf_draft.txt" in text
    assert "completed" in text
    assert "Analyzed words" in text
    assert "Coverage and limitations" in text
    assert "Matching passages" in text
    assert "pdf_reference.txt" in text
    assert "passageword0" in text
    assert "passageword1399" in text
    assert "[redacted]" in text
    assert "hunter2" not in text
    assert "QUETEXT_API_KEY" not in text

    from routers import plagiarism as plagiarism_router

    def fail_generation(*_args):
        raise ValueError("private renderer detail")

    monkeypatch.setattr(plagiarism_router, "build_plagiarism_pdf", fail_generation)
    failed = client.get(f"/api/plagiarism/reports/{scan_id}/pdf")
    assert failed.status_code == 500
    assert failed.json()["detail"] == "The PDF report could not be generated. Refresh the saved report and try again."
    assert "private renderer detail" not in failed.text


def test_plagiarism_local_comparison_excludes_self_and_other_users(monkeypatch):
    from fastapi.testclient import TestClient
    from app import app
    from services import plagiarism

    monkeypatch.setenv("QUETEXT_API_KEY", "must-not-be-used-for-local-scans")
    monkeypatch.setattr(plagiarism, "_quetext_request", lambda *_args, **_kwargs: pytest.fail("A local-only scan must not call an online API"))
    assert client.post("/api/register", json={"email": "plagiarism-scope@example.test", "password": "scope-test-123"}).status_code == 200
    passage = "This private passage provides enough distinct words to verify owner scoped plagiarism matching."
    reference = client.post(
        "/api/upload",
        files={"file": ("owner_reference.txt", passage.encode(), "text/plain")},
    )
    draft = client.post(
        "/api/upload",
        files={"file": ("owner_draft.txt", passage.encode(), "text/plain")},
    )
    assert reference.status_code == 200
    assert draft.status_code == 200

    other_client = TestClient(app)
    assert other_client.post("/api/register", json={"email": "plagiarism-other-scope@example.test", "password": "scope-test-456"}).status_code == 200
    other_upload = other_client.post(
        "/api/upload",
        files={"file": ("another_users_copy.txt", passage.encode(), "text/plain")},
    )
    assert other_upload.status_code == 200

    started = client.post("/api/plagiarism/scans", json={"file_id": draft.json()["id"]})
    report = client.get(f"/api/plagiarism/reports/{started.json()['scan_id']}").json()
    assert report["exact_match_count"] == 1
    assert [match["document_id"] for match in report["local_matches"]] == [reference.json()["id"]]
    assert report["local_matches"][0]["document_name"] == "owner_reference.txt"
    assert other_upload.json()["id"] not in {
        match["document_id"] for match in report["local_matches"]
    }


def test_plagiarism_supports_pdf_docx_txt_and_reports_scanned_pdf():
    from io import BytesIO
    from pypdf import PdfWriter

    assert client.post("/api/register", json={"email": "plagiarism-formats@example.test", "password": "formats-test-123"}).status_code == 200
    root = Path(__file__).resolve().parents[1] / "sample_documents"
    fixtures = [
        root / "pdf" / "02_sample_project_report.pdf",
        root / "docx" / "02_sample_project_documentation.docx",
        root / "txt" / "01_technical_notes.txt",
    ]
    for fixture in fixtures:
        response = client.post(
            "/api/upload",
            files={"file": (fixture.name, fixture.read_bytes(), "application/octet-stream")},
        )
        assert response.status_code == 200, response.text
        file_id = response.json()["id"]
        details = client.get(f"/api/files/{file_id}").json()
        assert details["has_text"] is True
        started = client.post("/api/plagiarism/scans", json={"file_id": file_id})
        assert started.status_code == 202, started.text
        report = client.get(f"/api/plagiarism/reports/{started.json()['scan_id']}").json()
        assert report["status"] in {"completed", "completed_with_warnings"}
        assert report["eligible_words_analyzed"] > 0

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    scanned_pdf = BytesIO()
    writer.write(scanned_pdf)
    uploaded = client.post(
        "/api/upload",
        files={"file": ("scanned_document.pdf", scanned_pdf.getvalue(), "application/pdf")},
    )
    assert uploaded.status_code == 200
    started = client.post("/api/plagiarism/scans", json={"file_id": uploaded.json()["id"]})
    report = client.get(f"/api/plagiarism/reports/{started.json()['scan_id']}").json()
    assert report["errors"]
    assert "OCR is not configured" in report["errors"][0]
    assert uploaded.json()["text_extracted"] is False


def _fake_quetext_response(status_code, data=None, message=None):
    class FakeResponse:
        def __init__(self):
            self.status_code = status_code
        @staticmethod
        def json():
            return {"status": status_code < 400, "data": data or {}, "message": message}
    return FakeResponse()


def test_plagiarism_quetext_report_shows_real_passage_source_and_similarity(monkeypatch):
    from services import plagiarism

    api_key = "test-quetext-key-not-for-reports"
    monkeypatch.setenv("QUETEXT_API_KEY", api_key)
    monkeypatch.setattr(plagiarism.time, "sleep", lambda *_args: None)
    assert client.post("/api/register", json={"email": "plagiarism-quetext@example.test", "password": "quetext-test-123"}).status_code == 200
    passage = "A distinctive marine biology report describes careful measurements across protected coastal habitats during repeated seasonal sampling by university teams along the conservation shoreline."
    uploaded = client.post("/api/upload", files={"file": ("quetext.txt", passage.encode(), "text/plain")})
    requests = []

    def fake_request(method, path, key, json_body=None):
        requests.append((method, path, key, json_body))
        if method == "POST":
            return _fake_quetext_response(200, {"id": "report-123"})
        if path.endswith("report-progress/report-123"):
            return _fake_quetext_response(200, [{"Progress": 1.0, "id": "report-123"}])
        return _fake_quetext_response(200, {"status": "completed", "score": 91.2, "matches": [{
            "input_text_match": passage,
            "highlighted_snippet": "<b>Matching coastal field measurements</b> from the public report.",
            "percent_similar": 91.2,
            "source": {"title": "Coastal research report", "url": "https://research.example.org/coastal"},
        }]})

    monkeypatch.setattr(plagiarism, "_quetext_request", fake_request)
    started = client.post("/api/plagiarism/scans", json={"file_id": uploaded.json()["id"], "online_search": True})
    report = client.get(f"/api/plagiarism/reports/{started.json()['scan_id']}").json()
    assert report["status"] == "completed"
    assert report["online_search_provider"] == "quetext"
    assert report["online_search_status"] == "completed"
    assert report["local_matches"] == []
    assert report["matches"] == []
    assert report["exact_match_count"] is None
    assert report["online_matches"][0]["passage"] == passage
    match = report["online_matches"][0]["sources"][0]
    assert match["title"] == "Coastal research report"
    assert match["url"] == "https://research.example.org/coastal"
    assert match["evidence"] == "Matching coastal field measurements from the public report."
    assert match["similarity_percent"] == 91.2
    assert all(item[2] == api_key for item in requests)
    assert all(item[0] != "POST" or len(item[3]["text"].split()) >= 20 for item in requests)
    assert api_key not in str(report)
    repeated = client.post("/api/plagiarism/scans", json={"file_id": uploaded.json()["id"], "online_search": True})
    assert repeated.json()["scan_id"] == started.json()["scan_id"]
    assert sum(item[0] == "POST" for item in requests) == 1


def test_plagiarism_quetext_long_progress_is_persisted_without_resubmission(monkeypatch):
    from services import plagiarism

    monkeypatch.setenv("QUETEXT_API_KEY", "configured-test-key")
    monkeypatch.setattr(plagiarism.time, "sleep", lambda *_args: None)
    assert client.post("/api/register", json={"email": "plagiarism-quetext-slow@example.test", "password": "quetext-slow-123"}).status_code == 200
    passage = "This longer provider test contains enough words to validate that delayed Quetext progress continues and real results eventually display in the existing SmartFiles plagiarism result section."
    uploaded = client.post("/api/upload", files={"file": ("slow-quetext.txt", passage.encode(), "text/plain")})
    counts = {"post": 0, "progress": 0}

    def fake_request(method, path, *_args, **_kwargs):
        if method == "POST":
            counts["post"] += 1
            return _fake_quetext_response(200, {"id": "slow-report"})
        if "progress" in path:
            counts["progress"] += 1
            value = 0.25 if counts["progress"] < 65 else 1.0
            return _fake_quetext_response(200, [{"Progress": value}])
        return _fake_quetext_response(200, {"status": "completed", "score": 34.5, "matches": [{
            "input_text_match": passage,
            "highlighted_snippet": "An excerpt returned by the provider.",
            "percent_similar": 87.5,
            "source": {"title": "Published research", "url": "https://research.example.org/report"},
        }]})

    monkeypatch.setattr(plagiarism, "_quetext_request", fake_request)
    started = client.post("/api/plagiarism/scans", json={"file_id": uploaded.json()["id"], "online_search": True})
    assert started.status_code == 202
    report = client.get(f"/api/plagiarism/reports/{started.json()['scan_id']}").json()
    assert report["online_search_status"] == "completed"
    assert report["quetext_score_percent"] == 34.5
    assert report["quetext_progress_percent"] == 100
    assert report["online_matches"][0]["sources"][0]["url"] == "https://research.example.org/report"
    assert counts == {"post": 1, "progress": 65}
    assert report["updated_at"]


def test_plagiarism_quetext_retry_resumes_saved_report_id(monkeypatch):
    from services import plagiarism

    monkeypatch.setenv("QUETEXT_API_KEY", "configured-test-key")
    monkeypatch.setattr(plagiarism.time, "sleep", lambda *_args: None)
    assert client.post("/api/register", json={"email": "plagiarism-quetext-resume@example.test", "password": "quetext-resume-123"}).status_code == 200
    passage = "A saved provider report can be resumed after a temporary rate limit without sending the paid document again and the returned source remains real."
    uploaded = client.post("/api/upload", files={"file": ("resume-quetext.txt", passage.encode(), "text/plain")})
    calls = {"post": 0, "progress": 0}

    def fake_request(method, path, *_args, **_kwargs):
        if method == "POST":
            calls["post"] += 1
            return _fake_quetext_response(200, {"id": "resumable-report"})
        if "progress" in path:
            calls["progress"] += 1
            if calls["progress"] == 1:
                return _fake_quetext_response(429, message="Rate limit exceeded")
            return _fake_quetext_response(200, [{"Progress": 1.0}])
        return _fake_quetext_response(200, {"status": "completed", "score": 8.0, "matches": [{
            "input_text_match": passage, "highlighted_snippet": "Actual provider excerpt.",
            "percent_similar": 80, "source": {"url": "https://example.org/saved-report"},
        }]})

    monkeypatch.setattr(plagiarism, "_quetext_request", fake_request)
    first = client.post("/api/plagiarism/scans", json={"file_id": uploaded.json()["id"], "online_search": True})
    failed_report = client.get(f"/api/plagiarism/reports/{first.json()['scan_id']}").json()
    assert failed_report["online_search_status"] == "rate_limited"
    second = client.post("/api/plagiarism/scans", json={"file_id": uploaded.json()["id"], "online_search": True})
    resumed = client.get(f"/api/plagiarism/reports/{second.json()['scan_id']}").json()
    assert second.json()["scan_id"] == first.json()["scan_id"]
    assert resumed["online_search_status"] == "completed"
    assert resumed["online_matches"][0]["sources"][0]["url"] == "https://example.org/saved-report"
    assert calls["post"] == 1


def test_plagiarism_quetext_frontend_polls_without_fixed_timeout_and_shows_score():
    script = (Path(__file__).resolve().parents[1] / "static" / "js" / "plagiarism.js").read_text(encoding="utf-8")
    assert "while (true)" in script
    assert "attempt < 125" not in script
    assert "quetext_progress_percent" in script
    assert "Overall Quetext plagiarism score" in script
    assert "Quetext did not provide a separate source excerpt" in script
    assert "Topic passage from your document" in script
    assert 'source.verification !== "provider_match"' in script
    assert 'sources.set(match.url' in script


def test_plagiarism_quetext_credits_and_missing_key_are_clear(monkeypatch):
    from services import plagiarism

    assert client.post("/api/register", json={"email": "plagiarism-quetext-errors@example.test", "password": "quetext-errors-123"}).status_code == 200
    text = "A sufficiently long reference passage supports local matching while the remote provider reports an issue during the carefully planned automated test today."
    reference = client.post("/api/upload", files={"file": ("reference.txt", text.encode(), "text/plain")})
    draft = client.post("/api/upload", files={"file": ("draft.txt", text.encode(), "text/plain")})
    monkeypatch.delenv("QUETEXT_API_KEY", raising=False)
    monkeypatch.setattr(plagiarism, "_quetext_request", lambda *_args, **_kwargs: pytest.fail("No request without a key"))
    missing = client.post("/api/plagiarism/scans", json={"file_id": draft.json()["id"], "online_search": True})
    missing_report = client.get(f"/api/plagiarism/reports/{missing.json()['scan_id']}").json()
    assert missing_report["online_search_status"] == "unavailable_missing_api_key"
    assert any("QUETEXT_API_KEY" in item for item in missing_report["errors"])
    assert missing_report["exact_match_count"] == 0
    assert missing_report["local_matches"] == []

    monkeypatch.setenv("QUETEXT_API_KEY", "configured-test-key")
    monkeypatch.setattr(plagiarism, "_quetext_request", lambda *_args, **_kwargs: _fake_quetext_response(402, {"balance_words": 0}, "Insufficient API credits"))
    exhausted = client.post("/api/plagiarism/scans", json={"file_id": draft.json()["id"], "online_search": True})
    exhausted_report = client.get(f"/api/plagiarism/reports/{exhausted.json()['scan_id']}").json()
    assert exhausted_report["online_search_status"] == "quota_exhausted"
    assert any("credits" in item.casefold() for item in exhausted_report["errors"])
    assert exhausted_report["exact_match_count"] == 0
    assert exhausted_report["local_matches"] == []
    assert "local_fallback" not in exhausted_report
    assert exhausted_report["online_matches"] == []


def test_plagiarism_quetext_rate_limit_is_reported_without_fabricated_matches(monkeypatch):
    from services import plagiarism

    monkeypatch.setenv("QUETEXT_API_KEY", "configured-test-key")
    assert client.post("/api/register", json={"email": "plagiarism-quetext-rate@example.test", "password": "quetext-rate-123"}).status_code == 200
    text = "This distinct passage contains enough words for checking the Quetext rate limit and local comparison behavior during a carefully controlled automated test today."
    uploaded = client.post("/api/upload", files={"file": ("rate.txt", text.encode(), "text/plain")})
    monkeypatch.setattr(plagiarism, "_quetext_request", lambda *_args, **_kwargs: _fake_quetext_response(429, message="Rate limit exceeded"))
    started = client.post("/api/plagiarism/scans", json={"file_id": uploaded.json()["id"], "online_search": True})
    report = client.get(f"/api/plagiarism/reports/{started.json()['scan_id']}").json()
    assert report["online_search_status"] == "rate_limited"
    assert report["online_matches"] == []
    assert report["sources"] == []
    assert len(report["errors"]) == len(set(report["errors"]))


@pytest.mark.parametrize("filename,content_type", [
    ("quetext.txt", "text/plain"), ("quetext.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ("quetext.pdf", "application/pdf"),
])
def test_plagiarism_quetext_processes_supported_text_documents(monkeypatch, filename, content_type):
    from io import BytesIO
    from docx import Document
    from reportlab.pdfgen import canvas
    from services import plagiarism

    monkeypatch.setenv("QUETEXT_API_KEY", "configured-test-key")
    monkeypatch.setattr(plagiarism.time, "sleep", lambda *_args: None)
    passage = "A distinctive document passage includes enough words to demonstrate the Quetext reporting workflow." \
        " It reports actual matches returned for source pages without making up websites or copied content during this integration test."
    if filename.endswith(".docx"):
        doc = Document(); doc.add_paragraph(passage); buffer = BytesIO(); doc.save(buffer); data = buffer.getvalue()
    elif filename.endswith(".pdf"):
        buffer = BytesIO(); pdf = canvas.Canvas(buffer); pdf.setFont("Helvetica", 8); pdf.drawString(20, 700, passage); pdf.save(); data = buffer.getvalue()
    else:
        data = passage.encode()
    uploaded = client.post("/api/upload", files={"file": (filename, data, content_type)})

    def fake_request(method, path, *_args, **_kwargs):
        if method == "POST": return _fake_quetext_response(200, {"id": "format-report"})
        if "progress" in path: return _fake_quetext_response(200, [{"Progress": 1}])
        return _fake_quetext_response(200, {"matches": [{"input_text_match": passage, "highlighted_snippet": "Actual returned source excerpt.", "percent_similar": 88, "source": {"url": "https://example.org/source"}}]})

    monkeypatch.setattr(plagiarism, "_quetext_request", fake_request)
    started = client.post("/api/plagiarism/scans", json={"file_id": uploaded.json()["id"], "online_search": True})
    report = client.get(f"/api/plagiarism/reports/{started.json()['scan_id']}").json()
    assert report["online_search_status"] == "completed"
    assert report["online_matches"][0]["sources"][0]["matched_text"] == passage
    assert report["online_matches"][0]["sources"][0]["url"] == "https://example.org/source"

def test_plagiarism_labels_fictional_demo_reports_and_rejects_unsupported_formats():
    from io import BytesIO
    from PIL import Image

    assert client.post("/api/register", json={"email": "plagiarism-demo@example.test", "password": "demo-test-123"}).status_code == 200
    imported = client.post("/api/demo/import")
    assert imported.status_code == 200, imported.text
    sample = next(
        file for file in client.get("/api/files").json()
        if file["is_demo"] and file["name"].lower().endswith(".txt")
    )
    started = client.post("/api/plagiarism/scans", json={"file_id": sample["id"]})
    report = client.get(f"/api/plagiarism/reports/{started.json()['scan_id']}").json()
    assert report["sample_data"] is True
    history = client.get("/api/plagiarism/reports").json()
    assert history[0]["sample_data"] is True

    image = BytesIO()
    Image.new("RGB", (2, 2), "white").save(image, format="PNG")
    uploaded = client.post("/api/upload", files={"file": ("not_supported.png", image.getvalue(), "image/png")})
    rejected = client.post("/api/plagiarism/scans", json={"file_id": uploaded.json()["id"]})
    assert rejected.status_code == 415


def test_plagiarism_reports_are_authenticated_and_owner_scoped():
    from fastapi.testclient import TestClient
    from app import app

    assert client.post("/api/register", json={"email": "plagiarism-owner@example.test", "password": "owner-test-123"}).status_code == 200
    uploaded = client.post(
        "/api/upload",
        files={"file": ("private_notes.txt", b"Private source comparison requires a sufficiently long passage.", "text/plain")},
    )
    started = client.post("/api/plagiarism/scans", json={"file_id": uploaded.json()["id"]})
    report_id = started.json()["scan_id"]
    second_report_id = client.post("/api/plagiarism/scans", json={"file_id": uploaded.json()["id"]}).json()["scan_id"]

    other_client = TestClient(app)
    assert other_client.post("/api/register", json={"email": "plagiarism-other@example.test", "password": "other-test-123"}).status_code == 200
    assert other_client.get(f"/api/plagiarism/reports/{report_id}").status_code == 404
    assert other_client.get(f"/api/plagiarism/reports/{report_id}/pdf").status_code == 404
    assert other_client.get(f"/api/plagiarism/comparisons/{report_id}/{second_report_id}").status_code == 404
    assert other_client.get("/api/plagiarism/reports", params={"file_id": uploaded.json()["id"]}).status_code == 404
    assert other_client.post("/api/plagiarism/scans", json={"file_id": uploaded.json()["id"]}).status_code == 404
    assert other_client.get("/api/plagiarism/reports").json() == []

    client.post("/api/logout")
    assert client.get("/api/plagiarism/reports").status_code == 401
    assert client.get(f"/api/plagiarism/reports/{report_id}/pdf").status_code == 401
    assert client.get(f"/api/plagiarism/comparisons/{report_id}/{second_report_id}").status_code == 401


def test_plagiarism_report_survives_reopening_the_sqlite_database():
    from database import connection, repository

    assert client.post("/api/register", json={"email": "plagiarism-restart@example.test", "password": "restart-test-123"}).status_code == 200
    uploaded = client.post(
        "/api/upload",
        files={"file": ("restart_notes.txt", b"A persisted report must remain available after the database connection restarts.", "text/plain")},
    )
    started = client.post("/api/plagiarism/scans", json={"file_id": uploaded.json()["id"]})
    report_id = started.json()["scan_id"]

    connection.DB.close()
    reopened = sqlite3.connect(str(connection.DB_PATH), check_same_thread=False)
    reopened.row_factory = sqlite3.Row
    connection.DB = reopened
    repository.DB = reopened

    restored = client.get(f"/api/plagiarism/reports/{report_id}")
    assert restored.status_code == 200
    assert restored.json()["document_name"] == "restart_notes.txt"
    assert restored.json()["status"] == "completed"
